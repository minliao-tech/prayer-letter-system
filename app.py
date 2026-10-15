from flask import Flask, request, redirect, url_for, render_template, send_file, flash
from pathlib import Path
from datetime import date, datetime, timedelta
import sqlite3
import os
import re
import secrets
import uuid
import io

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


app = Flask(__name__)
app.secret_key = "change-this-before-production"
app.json.sort_keys = False

BASE = Path(__file__).parent
DB = BASE / "prayer_letters.db"


# =========================================================
# 組別與單位
# =========================================================

GROUPS = {
    "A": {
        "主任牧師室": [],
        "聖工聯席會": [],
        "人資財會處": [],
        "行政資訊處": [],

        "創意藝術媒體處": [
            "媒體製作部",
            "雲端教會部",
            "企劃部",
            "崇拜部",
            "藝術部"
        ],

        "宣教植堂處": [],

        "啟示性事奉處": [
            "先知性事工中心",
            "禱告中心"
        ],

        "靈糧全球使徒性網絡": []
    },

    "B": {
        "社會服務處": [
            "多關懷服務部",
            "婦女培力部",
            "社區家庭關懷部",
            "福音外展部"
        ],

        "愛鄰協會": [],
        "事業處": [],
        "創新育成中心": []
    },

    "C": {
        "牧養總部": [
            "成人牧區",
            "北區會堂",
            "兒童牧區",
            "天使心牧區",
            "聯合崇拜",
            "學生牧區",
            "職場牧區",
            "喜樂家族",
            "台語牧區",
            "客語牧區"
        ],

        "國際事工中心": [
            "英語牧區",
            "印尼牧區",
            "越南牧區",
            "菲律賓牧區"
        ],

        "牧養支援處": [
            "牧養企劃部",
            "教育訓練部",
            "KC館"
        ],

        "基層福音事工處": [],

        "全人關顧中心": [],
        
        "靈糧國度領袖學院": [
            "神學院",
            "生命培訓學院",
            "巴拿巴宣教學院",
            "職場轉化學院"
        ]        
    },

    "D": {
        "福音中心": [
            "金門",
            "興隆",
            "萬金",
            "古亭",
            "澎湖",
            "竹圍",
            "三樹",
            "五股",
            "萬華",
            "東石"
        ]
    }
}


# =========================================================
# 輪值設定
# =========================================================

ANCHOR_DATE = date(2026, 9, 25)

ORDER = [
    "A",
    "B",
    "C",
    "D"
]

SKIP_DATES = {
    date(2026, 10, 9),   # 國慶連假
}


# =========================================================
# 資料庫
# =========================================================

SQLITE_SCHEMA = "\n    CREATE TABLE IF NOT EXISTS submissions(\n        id INTEGER PRIMARY KEY AUTOINCREMENT,\n        week_date TEXT NOT NULL,\n        group_code TEXT NOT NULL,\n        parent_unit TEXT NOT NULL,\n        unit_name TEXT NOT NULL,\n        contact_name TEXT,\n        contact_email TEXT,\n        status TEXT NOT NULL DEFAULT 'submitted',\n        created_at TEXT NOT NULL\n    );\n\n    CREATE TABLE IF NOT EXISTS prayer_items(\n        id INTEGER PRIMARY KEY AUTOINCREMENT,\n        submission_id INTEGER NOT NULL,\n        sort_order INTEGER NOT NULL,\n        content TEXT NOT NULL,\n        FOREIGN KEY(submission_id)\n            REFERENCES submissions(id)\n            ON DELETE CASCADE\n    );\n\n    CREATE TABLE IF NOT EXISTS weekly_settings(\n        id INTEGER PRIMARY KEY AUTOINCREMENT,\n        week_date TEXT NOT NULL,\n        group_code TEXT NOT NULL,\n\n        meeting_date TEXT DEFAULT '',\n        pastors TEXT DEFAULT '',\n        special_event TEXT DEFAULT '',\n\n        updated_at TEXT NOT NULL,\n\n        UNIQUE(week_date, group_code)\n    );\n    "

SHEET_ID = os.environ.get(
    "GOOGLE_SHEET_ID", "162bNAUvU5RnJm99NmKSi4HZ1o7eE9Ml8lEMUj0aVTmM"
).strip()
SERVICE_ACCOUNT_FILE = os.environ.get(
    "GOOGLE_SERVICE_ACCOUNT_FILE", "/etc/secrets/google-service-account.json"
)

# Each change appends a revision. Old rows are retained for recovery.
# A revision's submission and prayer rows are written in one atomic batch.
SHEET_TABLES = {
    "submissions": ("繳交紀錄", [
        "id", "week_date", "group_code", "parent_unit", "unit_name",
        "contact_name", "contact_email", "status", "created_at", "revision"
    ]),
    "prayer_items": ("代禱事項", [
        "id", "submission_id", "sort_order", "content", "revision"
    ]),
    "weekly_settings": ("本週設定", [
        "id", "week_date", "group_code", "meeting_date", "pastors",
        "special_event", "updated_at", "revision"
    ])
}


def sheets_client():
    from google.oauth2.service_account import Credentials
    from googleapiclient.discovery import build

    if not Path(SERVICE_ACCOUNT_FILE).is_file():
        raise RuntimeError(
            "Google Sheets credentials missing. Add google-service-account.json "
            "as a Render Secret File and share the spreadsheet with its client_email."
        )
    credentials = Credentials.from_service_account_file(
        SERVICE_ACCOUNT_FILE,
        scopes=["https://www.googleapis.com/auth/spreadsheets"]
    )
    return build("sheets", "v4", credentials=credentials, cache_discovery=False)


def sheet_row(values):
    # stringValue preserves long integer IDs and user text beginning with '='.
    return {"values": [
        {"userEnteredValue": {"stringValue": "" if v is None else str(v)}}
        for v in values
    ]}


def ensure_sheet_layout(api):
    metadata = api.spreadsheets().get(
        spreadsheetId=SHEET_ID, fields="sheets.properties"
    ).execute(num_retries=2)
    existing = {s["properties"]["title"]: s["properties"]["sheetId"]
                for s in metadata.get("sheets", [])}
    requests = []
    for title, headers in SHEET_TABLES.values():
        if title not in existing:
            sheet_id = secrets.randbelow(2**30) + 1
            while sheet_id in existing.values():
                sheet_id = secrets.randbelow(2**30) + 1
            existing[title] = sheet_id
            requests.extend([
                {"addSheet": {"properties": {
                    "sheetId": sheet_id, "title": title,
                    "gridProperties": {"frozenRowCount": 1,
                                       "columnCount": len(headers), "rowCount": 1000}
                }}},
                {"appendCells": {"sheetId": sheet_id,
                                 "rows": [sheet_row(headers)],
                                 "fields": "userEnteredValue"}},
                {"repeatCell": {
                    "range": {"sheetId": sheet_id, "startRowIndex": 0, "endRowIndex": 1},
                    "cell": {"userEnteredFormat": {
                        "backgroundColor": {"red": 0.93, "green": 0.93, "blue": 0.93},
                        "textFormat": {"bold": True, "foregroundColor": {
                            "red": 0, "green": 0, "blue": 0}}
                    }}, "fields": "userEnteredFormat"
                }}
            ])
    if requests:
        # Another worker may create these tabs at the same time. Only accept
        # the failure if a fresh read confirms all expected tabs now exist.
        from googleapiclient.errors import HttpError
        try:
            api.spreadsheets().batchUpdate(
                spreadsheetId=SHEET_ID, body={"requests": requests}
            ).execute()
        except HttpError as error:
            fresh = api.spreadsheets().get(
                spreadsheetId=SHEET_ID, fields="sheets.properties"
            ).execute(num_retries=2)
            actual = {s["properties"]["title"]: s["properties"]["sheetId"]
                      for s in fresh.get("sheets", [])}
            if not all(title in actual for title, _ in SHEET_TABLES.values()):
                raise error
            existing = actual
    return {title: existing[title] for title, _ in SHEET_TABLES.values()}


class SheetsConnection:
    """A disposable SQL view of durable Google Sheets records.

    SQL runs in memory for compatibility with existing routes; successful
    commit means Google accepted the changes, not merely local SQLite.
    """

    def __init__(self, api, sheet_ids):
        self.api = api
        self.sheet_ids = sheet_ids
        self.connection = sqlite3.connect(":memory:")
        self.connection.row_factory = sqlite3.Row
        try:
            self.connection.executescript(SQLITE_SCHEMA)
            self._load()
            self.original_submissions = self._submissions()
            self.original_settings = self._settings()
        except Exception:
            self.connection.close()
            raise

    def _load(self):
        result = self.api.spreadsheets().values().batchGet(
            spreadsheetId=SHEET_ID,
            ranges=[f"'{title}'!A:{chr(64 + len(headers))}"
                    for title, headers in SHEET_TABLES.values()],
            valueRenderOption="UNFORMATTED_VALUE"
        ).execute(num_retries=2)
        ranges = result.get("valueRanges", [])
        if len(ranges) != len(SHEET_TABLES):
            raise RuntimeError("Incomplete Google Sheets response; no data was written.")
        records = {}
        for (table, (title, headers)), data in zip(SHEET_TABLES.items(), ranges):
            rows = data.get("values", [])
            if not rows or rows[0] != headers:
                raise RuntimeError(f"Header mismatch in {title}; restore the original headers.")
            records[table] = []
            for values in rows[1:]:
                if not any(str(value) for value in values):
                    continue
                row = dict(zip(headers, list(values) + [""] * (len(headers) - len(values))))
                if not row.get("id") or not row.get("revision"):
                    raise RuntimeError(f"Incomplete record in {title}; check the sheet before continuing.")
                for field in ("id", "submission_id", "sort_order"):
                    if field in row:
                        row[field] = int(row[field])
                records[table].append(row)

        latest_submissions = {}
        for row in records["submissions"]:
            latest_submissions[row["id"]] = row
        latest_settings = {}
        for row in records["weekly_settings"]:
            latest_settings[(row["week_date"], row["group_code"])] = row

        def insert(table, row):
            columns = SHEET_TABLES[table][1][:-1]
            self.connection.execute(
                f"INSERT INTO {table} ({','.join(columns)}) "
                f"VALUES ({','.join('?' for _ in columns)})",
                [row[c] for c in columns]
            )

        for row in latest_submissions.values():
            insert("submissions", row)
        for row in latest_settings.values():
            insert("weekly_settings", row)
        seen = set()
        for row in records["prayer_items"]:
            sub = latest_submissions.get(row["submission_id"])
            if sub and row["revision"] == sub["revision"] and row["id"] not in seen:
                insert("prayer_items", row)
                seen.add(row["id"])
        self.connection.commit()

    def execute(self, sql, params=()):
        # Random IDs prevent two simultaneous submissions choosing max(id)+1.
        match = re.match(r"\s*INSERT\s+INTO\s+(submissions|prayer_items|weekly_settings)\s*\(",
                         sql, re.IGNORECASE)
        if match:
            table = match.group(1).lower()
            new_id = secrets.randbelow(2**62 - 1) + 1
            while self.connection.execute(
                f"SELECT id FROM {table} WHERE id=?", (new_id,)
            ).fetchone():
                new_id = secrets.randbelow(2**62 - 1) + 1
            sql = sql[:match.end()] + "id, " + sql[match.end():]
            sql = re.sub(r"\bVALUES\s*\(", "VALUES(?, ", sql, count=1, flags=re.IGNORECASE)
            params = (new_id,) + tuple(params)
        return self.connection.execute(sql, params)

    def executescript(self, sql):
        return self.connection.executescript(sql)

    def _submissions(self):
        result = {}
        for row in self.connection.execute("SELECT * FROM submissions"):
            items = [dict(item) for item in self.connection.execute(
                "SELECT * FROM prayer_items WHERE submission_id=? ORDER BY sort_order, id",
                (row["id"],)
            )]
            result[row["id"]] = (dict(row), items)
        return result

    def _settings(self):
        return {(row["week_date"], row["group_code"]): dict(row)
                for row in self.connection.execute("SELECT * FROM weekly_settings")}

    def commit(self):
        submissions = self._submissions()
        settings = self._settings()
        pending = {table: [] for table in SHEET_TABLES}
        revision = uuid.uuid4().hex
        for sid, (sub, items) in submissions.items():
            if self.original_submissions.get(sid) != (sub, items):
                pending["submissions"].append(
                    [sub[c] for c in SHEET_TABLES["submissions"][1][:-1]] + [revision]
                )
                for item in items:
                    pending["prayer_items"].append(
                        [item[c] for c in SHEET_TABLES["prayer_items"][1][:-1]] + [revision]
                    )
        for key, row in settings.items():
            if self.original_settings.get(key) != row:
                pending["weekly_settings"].append(
                    [row[c] for c in SHEET_TABLES["weekly_settings"][1][:-1]] + [revision]
                )
        requests = []
        for table, rows in pending.items():
            if rows:
                title = SHEET_TABLES[table][0]
                requests.append({"appendCells": {
                    "sheetId": self.sheet_ids[title],
                    "rows": [sheet_row(row) for row in rows],
                    "fields": "userEnteredValue"
                }})
        if requests:
            # Submissions and their items are committed atomically. Do not
            # return a success page when this call fails or times out.
            self.api.spreadsheets().batchUpdate(
                spreadsheetId=SHEET_ID, body={"requests": requests}
            ).execute(num_retries=2)
        self.connection.commit()
        self.original_submissions = submissions
        self.original_settings = settings

    def close(self):
        self.connection.close()
        http = getattr(self.api, "_http", None)
        if hasattr(http, "close"):
            http.close()


_sheet_ids = None


def db():
    global _sheet_ids
    if not SHEET_ID:
        raise RuntimeError("GOOGLE_SHEET_ID must not be empty.")
    api = sheets_client()
    if _sheet_ids is None:
        _sheet_ids = ensure_sheet_layout(api)
    return SheetsConnection(api, _sheet_ids)


@app.errorhandler(Exception)
def storage_error(error):
    from werkzeug.exceptions import HTTPException
    if isinstance(error, HTTPException):
        return error
    app.logger.error("Request failed (%s).", type(error).__name__)
    return (
        "資料讀取或儲存未能確認成功，請保留您填寫的內容並稍後重試。"
        "若正在提交資料，請先查看後台是否已有紀錄，避免重複繳交。"
        "請管理員檢查 Google Sheets 連線設定。",
        503
    )



def init_db():

    con = db()

    con.executescript("""
    CREATE TABLE IF NOT EXISTS submissions(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        week_date TEXT NOT NULL,
        group_code TEXT NOT NULL,
        parent_unit TEXT NOT NULL,
        unit_name TEXT NOT NULL,
        contact_name TEXT,
        contact_email TEXT,
        status TEXT NOT NULL DEFAULT 'submitted',
        created_at TEXT NOT NULL
    );

    CREATE TABLE IF NOT EXISTS prayer_items(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        submission_id INTEGER NOT NULL,
        sort_order INTEGER NOT NULL,
        content TEXT NOT NULL,
        FOREIGN KEY(submission_id)
            REFERENCES submissions(id)
            ON DELETE CASCADE
    );

    CREATE TABLE IF NOT EXISTS weekly_settings(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        week_date TEXT NOT NULL,
        group_code TEXT NOT NULL,

        meeting_date TEXT DEFAULT '',
        pastors TEXT DEFAULT '',
        special_event TEXT DEFAULT '',

        updated_at TEXT NOT NULL,

        UNIQUE(week_date, group_code)
    );
    """)

    con.commit()
    con.close()


# =========================================================
# 星期五
# =========================================================

def friday_of_week(d=None):

    d = d or date.today()

    if d.weekday() <= 4:

        return d + timedelta(
            days=(4 - d.weekday())
        )

    return d - timedelta(
        days=d.weekday() - 4
    )


def is_skip_week(week_date):

    return week_date in SKIP_DATES


# =========================================================
# 計算輪值
# =========================================================

def group_for_week(week_date):

    if is_skip_week(week_date):
        return None

    if week_date < ANCHOR_DATE:

        valid_weeks = 0

        cursor = (
            ANCHOR_DATE -
            timedelta(days=7)
        )

        while cursor >= week_date:

            if cursor not in SKIP_DATES:
                valid_weeks += 1

            cursor -= timedelta(days=7)

        return ORDER[
            (-valid_weeks) % len(ORDER)
        ]

    valid_weeks = 0

    cursor = ANCHOR_DATE

    while cursor < week_date:

        cursor += timedelta(days=7)

        if cursor > week_date:
            break

        if cursor not in SKIP_DATES:
            valid_weeks += 1

    return ORDER[
        valid_weeks % len(ORDER)
    ]


def next_active_friday(d=None):

    wd = friday_of_week(d)

    while wd in SKIP_DATES:
        wd += timedelta(days=7)

    return wd


def current_cycle():

    wd = next_active_friday()

    group_code = group_for_week(wd)

    return wd, group_code


# =========================================================
# 應繳單位
# =========================================================

def expected_units(group_code):

    out = []

    if not group_code:
        return out

    for parent, children in GROUPS[group_code].items():

        if children:

            for child in children:

                out.append(
                    (parent, child)
                )

        else:

            out.append(
                (parent, parent)
            )

    return out


# =========================================================
# 首頁
# =========================================================

@app.route("/")
def index():

    wd, group_code = current_cycle()

    return render_template(
        "index.html",
        week_date=wd.isoformat(),
        group_code=group_code,
        groups=GROUPS
    )


# =========================================================
# 填寫代禱事項
# =========================================================

@app.route("/submit", methods=["GET", "POST"])
def submit():

    default_week, default_group = current_cycle()

    if request.method == "POST":

        week_date = request.form["week_date"]
        group_code = request.form["group_code"]
        parent_unit = request.form["parent_unit"]
        unit_name = request.form["unit_name"]

        contact_name = request.form.get(
            "contact_name",
            ""
        ).strip()

        contact_email = request.form.get(
            "contact_email",
            ""
        ).strip()

        items = [
            x.strip()
            for x in request.form.getlist(
                "prayer_item"
            )
            if x.strip()
        ]

        if not items:

            flash(
                "請至少填寫一項代禱事項。"
            )

            return redirect(
                url_for("submit")
            )

        con = db()

        cur = con.execute(
            """
            INSERT INTO submissions
            (
                week_date,
                group_code,
                parent_unit,
                unit_name,
                contact_name,
                contact_email,
                status,
                created_at
            )
            VALUES(?,?,?,?,?,?,?,?)
            """,
            (
                week_date,
                group_code,
                parent_unit,
                unit_name,
                contact_name,
                contact_email,
                "submitted",
                datetime.now().isoformat(
                    timespec="microseconds"
                )
            )
        )

        sid = cur.lastrowid

        for i, item in enumerate(
            items,
            1
        ):

            con.execute(
                """
                INSERT INTO prayer_items
                (
                    submission_id,
                    sort_order,
                    content
                )
                VALUES(?,?,?)
                """,
                (
                    sid,
                    i,
                    item
                )
            )

        con.commit()
        con.close()

        return render_template(
            "thanks.html",
            unit_name=unit_name,
            count=len(items)
        )

    return render_template(
        "submit.html",
        week_date=default_week.isoformat(),
        group_code=default_group,
        groups=GROUPS
    )


# =========================================================
# 單位 API
# =========================================================

@app.route("/api/units/<group_code>")
def units(group_code):

    return GROUPS.get(
        group_code,
        {}
    )


# =========================================================
# 管理後台
# =========================================================

@app.route("/admin")
def admin():

    week = request.args.get("week")

    if week:

        wd = date.fromisoformat(
            week
        )

    else:

        wd, _ = current_cycle()

    requested_group = request.args.get(
        "group"
    )

    group_code = (
        requested_group
        or group_for_week(wd)
    )

    if not group_code:

        wd = next_active_friday(
            wd + timedelta(days=1)
        )

        group_code = group_for_week(wd)

    expected = expected_units(
        group_code
    )

    con = db()

    rows = con.execute(
        """
        SELECT
            s.*,

            (
                SELECT COUNT(*)
                FROM prayer_items p
                WHERE p.submission_id=s.id
            ) item_count

        FROM submissions s

        WHERE
            week_date=?
            AND group_code=?

        ORDER BY created_at DESC
        """,
        (
            wd.isoformat(),
            group_code
        )
    ).fetchall()

    latest = {}

    for r in rows:

        key = (
            r["parent_unit"],
            r["unit_name"]
        )

        if key not in latest:
            latest[key] = r

    status_rows = []

    for parent, unit in expected:

        status_rows.append({
            "parent": parent,
            "unit": unit,
            "submission":
                latest.get(
                    (parent, unit)
                )
        })


    # 讀取本週 Word 前言設定

    weekly = con.execute(
        """
        SELECT *
        FROM weekly_settings
        WHERE week_date=? AND group_code=?
        """,
        (
            wd.isoformat(),
            group_code
        )
    ).fetchone()


    con.close()


    return render_template(
        "admin.html",

        week_date=wd.isoformat(),
        group_code=group_code,

        status_rows=status_rows,

        submitted=sum(
            1
            for x in status_rows
            if x["submission"]
        ),

        total=len(status_rows),

        weekly=weekly
    )


# =========================================================
# 儲存 Word 本週前言
# =========================================================

@app.route(
    "/admin/weekly-settings",
    methods=["POST"]
)
def save_weekly_settings():

    week_date = request.form[
        "week_date"
    ]

    group_code = request.form[
        "group_code"
    ]

    meeting_date = request.form.get(
        "meeting_date",
        ""
    ).strip()

    pastors = request.form.get(
        "pastors",
        ""
    ).strip()

    special_event = request.form.get(
        "special_event",
        ""
    ).strip()

    con = db()

    con.execute(
        """
        INSERT INTO weekly_settings
        (
            week_date,
            group_code,
            meeting_date,
            pastors,
            special_event,
            updated_at
        )

        VALUES(?,?,?,?,?,?)

        ON CONFLICT(week_date, group_code)

        DO UPDATE SET
            meeting_date=excluded.meeting_date,
            pastors=excluded.pastors,
            special_event=excluded.special_event,
            updated_at=excluded.updated_at
        """,
        (
            week_date,
            group_code,
            meeting_date,
            pastors,
            special_event,
            datetime.now().isoformat(
                timespec="microseconds"
            )
        )
    )

    con.commit()
    con.close()

    flash(
        "本週 Word 前言已儲存。"
    )

    return redirect(
        url_for(
            "admin",
            week=week_date,
            group=group_code
        )
    )


# =========================================================
# 後台修改代禱事項
# =========================================================

@app.route(
    "/admin/edit/<int:sid>",
    methods=["GET", "POST"]
)
def edit_submission(sid):

    con = db()

    sub = con.execute(
        """
        SELECT *
        FROM submissions
        WHERE id=?
        """,
        (sid,)
    ).fetchone()

    if not sub:

        con.close()

        return "Not found", 404

    if request.method == "POST":

        items = [
            x.strip()
            for x in request.form.getlist(
                "prayer_item"
            )
            if x.strip()
        ]

        con.execute(
            """
            UPDATE submissions
            SET status=?
            WHERE id=?
            """,
            (
                request.form.get(
                    "status",
                    "approved"
                ),
                sid
            )
        )

        con.execute(
            """
            DELETE FROM prayer_items
            WHERE submission_id=?
            """,
            (sid,)
        )

        for i, item in enumerate(
            items,
            1
        ):

            con.execute(
                """
                INSERT INTO prayer_items
                (
                    submission_id,
                    sort_order,
                    content
                )
                VALUES(?,?,?)
                """,
                (
                    sid,
                    i,
                    item
                )
            )

        con.commit()
        con.close()

        return redirect(
            url_for(
                "admin",
                week=sub["week_date"],
                group=sub["group_code"]
            )
        )

    items = con.execute(
        """
        SELECT *
        FROM prayer_items
        WHERE submission_id=?
        ORDER BY sort_order
        """,
        (sid,)
    ).fetchall()

    con.close()

    return render_template(
        "edit.html",
        sub=sub,
        items=items
    )


# =========================================================
# Word 字型
# =========================================================

def set_run_font(
    run,
    size=12,
    bold=False,
    color=None
):

    run.font.name = (
        "Microsoft JhengHei"
    )

    run._element.get_or_add_rPr()

    rFonts = (
        run._element.rPr
        .get_or_add_rFonts()
    )

    rFonts.set(
        qn("w:eastAsia"),
        "Microsoft JhengHei"
    )

    rFonts.set(
        qn("w:ascii"),
        "Microsoft JhengHei"
    )

    rFonts.set(
        qn("w:hAnsi"),
        "Microsoft JhengHei"
    )

    rFonts.set(
        qn("w:cs"),
        "Microsoft JhengHei"
    )

    run.font.size = Pt(size)
    run.bold = bold

    if color:

        run.font.color.rgb = (
            RGBColor(*color)
        )


# =========================================================
# Word 數字編號：換行文字對齊正文起點
# =========================================================

def set_numbered_paragraph(paragraph):

    paragraph.paragraph_format.left_indent = Cm(0.65)
    paragraph.paragraph_format.first_line_indent = Cm(-0.65)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(0.65))


# =========================================================
# Word 水平分隔線：設在現有段落下方，避免額外空白段落
# =========================================================

def add_horizontal_separator(paragraph):

    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))

    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)

    bottom = borders.find(qn("w:bottom"))

    if bottom is None:
        bottom = OxmlElement("w:bottom")
        borders.append(bottom)

    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), "000000")


# =========================================================
# Word 顏色
# =========================================================

PARENT_BLUE = (
    0,
    102,
    204
)

UNIT_GREEN = (
    0,
    100,
    0
)

PRAYER_BLACK = (
    0,
    0,
    0
)

HEADER_RED = (
    192,
    0,
    0
)

PASTOR_GREEN = (
    84,
    130,
    53
)


# =========================================================
# Word 第二層：
# 自動產生當週所有單位名稱
# =========================================================

def add_group_roster(
    doc,
    group_code
):

    for parent, children in (
        GROUPS[group_code].items()
    ):

        p = doc.add_paragraph()

        # 上層單位：藍色
        r = p.add_run(
            parent
        )

        set_run_font(
            r,
            12,
            True,
            PARENT_BLUE
        )

        if children:

            r = p.add_run("：")

            set_run_font(
                r,
                12,
                False,
                PRAYER_BLACK
            )

            for index, child in enumerate(
                children
            ):

                if index > 0:

                    r = p.add_run("、")

                    set_run_font(
                        r,
                        12,
                        False,
                        PRAYER_BLACK
                    )

                display_child = child

                # D組顯示「金門福音中心」等
                if (
                    group_code == "D"
                    and
                    parent == "福音中心"
                ):

                    display_child = (
                        f"{child}{parent}"
                    )

                r = p.add_run(
                    display_child
                )

                set_run_font(
                    r,
                    12,
                    False,
                    UNIT_GREEN
                )


# =========================================================
# 匯出 Word
# =========================================================

@app.route("/export")
def export_word():

    week = request.args.get(
        "week"
    )

    group_code = request.args.get(
        "group"
    )

    if not week or not group_code:

        return (
            "Missing week/group",
            400
        )

    con = db()


    # =====================================================
    # 取得本週前言
    # =====================================================

    weekly = con.execute(
        """
        SELECT *
        FROM weekly_settings
        WHERE week_date=? AND group_code=?
        """,
        (
            week,
            group_code
        )
    ).fetchone()


    # =====================================================
    # 取得代禱資料
    # =====================================================

    rows = con.execute(
        """
        SELECT *
        FROM submissions

        WHERE
            week_date=?
            AND group_code=?
            AND status IN (
                'submitted',
                'approved'
            )

        ORDER BY created_at, id
        """,
        (
            week,
            group_code
        )
    ).fetchall()


    latest = {}

    for r in rows:

        latest[
            (
                r["parent_unit"],
                r["unit_name"]
            )
        ] = r


    # =====================================================
    # 建立 Word
    # =====================================================

    doc = Document()

    # Word「窄」邊界：上、下、左、右各 1.27 公分
    for section in doc.sections:
        section.top_margin = Cm(1.27)
        section.bottom_margin = Cm(1.27)
        section.left_margin = Cm(1.27)
        section.right_margin = Cm(1.27)


    # =====================================================
    # 主標題
    # =====================================================

    title = doc.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    r = title.add_run(
        f"{week.replace('-', '')} "
        f"教會代禱信 "
        f"{group_code}組"
    )

    set_run_font(
        r,
        16,
        True,
        PRAYER_BLACK
    )


    # =====================================================
    # 第一層：本週聚會禱告
    # =====================================================

    p = doc.add_paragraph()

    r = p.add_run(
        "為本週主日崇拜、聚會、線上直播禱告："
    )

    set_run_font(
        r,
        12,
        True,
        PRAYER_BLACK
    )


    if weekly:

        meeting_date = (
            weekly["meeting_date"] or ""
        ).strip()

        pastors = (
            weekly["pastors"] or ""
        ).strip()

        special_event = (
            weekly["special_event"] or ""
        ).strip()


        # -------------------------------------------------
        # 1. 本週六日...
        # -------------------------------------------------

        if meeting_date or pastors:

            p = doc.add_paragraph()

            set_numbered_paragraph(p)

            r = p.add_run("1.\t")

            set_run_font(
                r,
                12,
                False,
                PRAYER_BLACK
            )


            if meeting_date:

                r = p.add_run(
                    f"本週六日({meeting_date})"
                    "在宣教、山莊、北區、各牧區、"
                    "各福音中心的講員"
                )

                set_run_font(
                    r,
                    12,
                    False,
                    PRAYER_BLACK
                )


            if pastors:

                r = p.add_run("（")

                set_run_font(
                    r,
                    12,
                    False,
                    PRAYER_BLACK
                )

                r = p.add_run(
                    pastors
                )

                set_run_font(
                    r,
                    12,
                    False,
                    PASTOR_GREEN
                )

                r = p.add_run(
                    "），以及敬拜團主領、主持人等"
                    "服事同工，並錄影連線、轉播上傳作業。"
                )

                set_run_font(
                    r,
                    12,
                    False,
                    PRAYER_BLACK
                )


        # -------------------------------------------------
        # a. 特別聚會
        # b. 特別聚會
        # -------------------------------------------------

        if special_event:

            p = doc.add_paragraph()

            r = p.add_run(
                f"a. {special_event}"
            )

            set_run_font(
                r,
                12,
                False,
                PRAYER_BLACK
            )


    # 第一層（包含特別聚會）結束後加入真正的段落下框線
    add_horizontal_separator(doc.paragraphs[-1])


    # =====================================================
    # 紅色標題
    # =====================================================

    p = doc.add_paragraph()

    r = p.add_run(
        "【為所提單位的同工禱告】"
    )

    set_run_font(
        r,
        12,
        True,
        HEADER_RED
    )


    # =====================================================
    # 第二層：
    # 自動列出當週輪值組所有單位
    # =====================================================

    add_group_roster(
        doc,
        group_code
    )


    # 當週所有單位名單結束後加入真正的段落下框線
    add_horizontal_separator(doc.paragraphs[-1])


    # =====================================================
    # 正式代禱內容
    # =====================================================

    for parent, children in (
        GROUPS[group_code].items()
    ):

        if children:

            has_submission = any(
                latest.get(
                    (parent, unit)
                )
                for unit in children
            )

            if not has_submission:
                continue


            # 上層單位
            p = doc.add_paragraph()

            r = p.add_run(
                parent
            )

            set_run_font(
                r,
                13,
                True,
                PARENT_BLUE
            )

            units = children

        else:

            units = [
                parent
            ]


        for unit in units:

            sub = latest.get(
                (
                    parent,
                    unit
                )
            )

            if not sub:
                continue


            if (
                group_code == "D"
                and
                parent == "福音中心"
            ):

                display = (
                    f"{unit}{parent}"
                )

            else:

                display = unit


            # 填寫單位
            p = doc.add_paragraph()

            r = p.add_run(
                f"【{display}】"
            )

            set_run_font(
                r,
                12,
                True,
                UNIT_GREEN
            )


            # 代禱事項
            items = con.execute(
                """
                SELECT *
                FROM prayer_items

                WHERE submission_id=?

                ORDER BY sort_order
                """,
                (
                    sub["id"],
                )
            ).fetchall()


            for i, item in enumerate(
                items,
                1
            ):

                p = doc.add_paragraph()

                set_numbered_paragraph(p)

                p.paragraph_format.space_after = (
                    Pt(4)
                )

                r = p.add_run(
                    f"{i}.\t{item['content']}"
                )

                set_run_font(
                    r,
                    12,
                    False,
                    PRAYER_BLACK
                )


    con.close()


    # =====================================================
    # 儲存 Word
    # =====================================================

    bio = io.BytesIO()

    doc.save(bio)

    bio.seek(0)


    filename = (
        f"{week.replace('-', '')} "
        f"教會代禱信 "
        f"{group_code}組.docx"
    )


    return send_file(
        bio,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            "application/vnd.openxmlformats-"
            "officedocument.wordprocessingml.document"
        )
    )


# =========================================================
# 初始化
# =========================================================

init_db()


# =========================================================
# 本機啟動
# =========================================================

if __name__ == "__main__":

    app.run(
        debug=True,
        port=5000
    )
